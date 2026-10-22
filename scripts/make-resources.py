#!/usr/bin/env python3
"""Builds the downloadable templates in public/resources/ (the Resources page links to them).

  docker run --rm -v "$PWD:/w" -w /w python:3.12-slim sh -c \
    'pip install -q python-docx openpyxl python-pptx reportlab && python scripts/make-resources.py'

Each file is a plain starting point, not legal or financial advice.
"""
import os

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "public", "resources")
os.makedirs(OUT, exist_ok=True)
NOTE = "This is a template to start from. It is not legal, tax or financial advice. Have a professional check anything that matters."


def doc(name, title, blocks):
    d = Document()
    d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)
    d.add_heading(title, 0)
    p = d.add_paragraph(NOTE); p.runs[0].italic = True
    for kind, text in blocks:
        if kind == "h": d.add_heading(text, 1)
        elif kind == "p": d.add_paragraph(text)
        elif kind == "b": d.add_paragraph(text, style="List Bullet")
        elif kind == "n": d.add_paragraph(text, style="List Number")
        elif kind == "t":
            t = d.add_table(rows=len(text), cols=len(text[0])); t.style = "Light Grid Accent 1"
            for i, row in enumerate(text):
                for j, cell in enumerate(row): t.cell(i, j).text = cell
            d.add_paragraph()
    d.save(os.path.join(OUT, name))


doc("founder-agreement.docx", "Founder Agreement (template, England and Wales)", [
    ("p", "Date: [date]. Between: [Founder A, address] and [Founder B, address] (the Founders), about the business known as [company or trading name] (the Business)."),
    ("h", "1. Purpose and roles"), ("p", "The Founders will build the Business together. Each Founder's role and time commitment:"),
    ("t", [["Founder", "Role", "Time commitment", "Expected start"], ["[Name]", "[e.g. product]", "[e.g. 20 hours a week]", "[date]"], ["[Name]", "[e.g. sales]", "[e.g. full time]", "[date]"]]),
    ("h", "2. Ownership"), ("p", "Ownership of the Business is split as follows, before any outside investment:"),
    ("t", [["Founder", "Percentage", "Reason for the split"], ["[Name]", "[ ]%", "[idea, cash, time, skills]"], ["[Name]", "[ ]%", "[ ]"]]),
    ("h", "3. Vesting"), ("p", "Each Founder's shares vest monthly over [4] years with a [12]-month cliff: nothing vests before the cliff, then 25% vests at the cliff and the rest monthly. Unvested shares are offered back to the Business at nominal value if a Founder leaves."),
    ("h", "4. Intellectual property"), ("p", "Each Founder assigns to the Business all work they create for it, including code, designs, brand, writing and data, and will sign any paper needed to confirm this. Work made before the Business started that is not listed here stays with its owner and is licensed to the Business: [list]."),
    ("h", "5. Decisions"), ("b", "Day-to-day decisions in a Founder's own area: that Founder decides."), ("b", "Decisions that need every Founder to agree: selling the Business, taking investment or debt over [amount], issuing shares, changing this agreement, spending over [amount]."),
    ("b", "If Founders deadlock for more than [30] days they will try mediation before anything else."),
    ("h", "6. Money"), ("p", "Salaries, expenses and loans from Founders: [describe]. Nobody is paid before the Business can afford it, except as agreed in writing."),
    ("h", "7. Leaving"), ("b", "Good leaver (illness, agreed departure): keeps vested shares, unvested shares return to the Business."),
    ("b", "Bad leaver (serious breach, leaving within the cliff, working for a competitor): the Business may buy back all shares at the lower of cost and fair value."),
    ("b", "A leaving Founder gives [3] months' notice where possible and helps with a handover."),
    ("h", "8. Confidentiality and competing work"), ("p", "Founders keep the Business's information confidential during and for [2] years after their involvement, and do not run a competing business while a Founder."),
    ("h", "9. Disputes and law"), ("p", "This agreement is governed by the law of England and Wales. Founders will try to settle disputes by talking, then mediation, before going to court."),
    ("h", "Signatures"), ("p", "Signed: ____________________   Date: ________        Signed: ____________________   Date: ________")])

doc("nda-template.docx", "Non-Disclosure Agreement (template, UK)", [
    ("p", "Choose one: ☐ One-way (only [Discloser] shares confidential information)   ☐ Mutual (both parties share confidential information)."),
    ("p", "Date: [date]. Parties: [Party A, address] and [Party B, address]."),
    ("h", "1. Purpose"), ("p", "The parties want to talk about [describe the purpose, for example a possible partnership] and may share confidential information to do so (the Purpose)."),
    ("h", "2. What is confidential"), ("p", "Anything shared for the Purpose that is marked confidential or that a reasonable person would see as confidential: plans, prices, customer lists, software, designs and know-how."),
    ("h", "3. What is not"), ("b", "Information already public, or that becomes public without the receiver's fault."), ("b", "Information the receiver already had, or got from someone else lawfully, or made up independently."),
    ("h", "4. Duties of the receiver"), ("b", "Keep it confidential and secure, using at least the care used for its own confidential information."), ("b", "Use it only for the Purpose."), ("b", "Share it only with people who need it for the Purpose and are bound to keep it confidential."),
    ("b", "Disclosure required by law or a court is allowed after telling the other party, where the law permits."),
    ("h", "5. Length"), ("p", "The duty to keep information confidential lasts [3] years from the date above. Trade secrets stay protected for as long as they remain secret."),
    ("h", "6. Return"), ("p", "On request the receiver returns or deletes confidential information and confirms in writing that it has done so."),
    ("h", "7. No other rights"), ("p", "Nothing here gives either party ownership of the other's information, or a duty to go ahead with any deal."),
    ("h", "8. Law"), ("p", "The law of England and Wales applies and its courts have jurisdiction."),
    ("p", "Signed: ____________________ for [Party A]   Date: ________        Signed: ____________________ for [Party B]   Date: ________")])

doc("investor-update.docx", "Monthly Investor Update (one page)", [
    ("p", "Month: [month year]   Company: [name]   Written by: [name]"),
    ("h", "Headline"), ("p", "One sentence on how the month went: [for example: revenue up 20%, signed our first three customers]."),
    ("h", "Highlights"), ("b", "[win 1]"), ("b", "[win 2]"), ("b", "[win 3]"),
    ("h", "Lowlights and what we learned"), ("b", "[what did not work, and what changed because of it]"),
    ("h", "Key numbers"), ("t", [["Metric", "This month", "Last month", "Change"], ["Revenue (£)", "", "", ""], ["Customers", "", "", ""], ["Cash in bank (£)", "", "", ""], ["Monthly cost (£)", "", "", ""], ["Months of cash left", "", "", ""]]),
    ("h", "Asks"), ("b", "[an introduction to ...]"), ("b", "[advice on ...]"),
    ("h", "Next month"), ("b", "[goal 1]"), ("b", "[goal 2]")])

doc("job-description.docx", "Early Hire Job Description (template)", [
    ("h", "[Job title] at [company]"), ("p", "Location: [place or remote]. Hours: [full time / part time]. Pay: £[range] a year plus [equity or other]. Reports to: [name]."),
    ("h", "About us"), ("p", "[Two or three sentences on what the business does, who it helps and where it is now.]"),
    ("h", "The role"), ("p", "[What this person will own in their first year, in plain words.]"),
    ("h", "In the first 90 days you will"), ("b", "[outcome 1]"), ("b", "[outcome 2]"), ("b", "[outcome 3]"),
    ("h", "You might be a good fit if you"), ("b", "[skill or experience that is needed]"), ("b", "[way of working that matters here]"), ("b", "[nice to have, clearly marked as optional]"),
    ("h", "What we offer"), ("b", "[pay and benefits]"), ("b", "[equity: a few lines on how it works and that it is not guaranteed to be worth anything]"), ("b", "[learning, flexibility]"),
    ("h", "How to apply"), ("p", "[how, by when, what the steps are]. We welcome applications from everyone and will make reasonable adjustments in the process. Note: check right-to-work and employment contract rules before hiring.")])

doc("one-page-plan.docx", "One-Page Business Plan", [
    ("p", "Fill in each box in plain words. If a box is hard to fill, that is the thing to find out first."),
    ("t", [["Section", "Your answer"],
           ["The problem: who has it, and what does it cost them today?", ""],
           ["Your solution: what exactly do you sell, with what is included and what is not?", ""],
           ["Your customer: one specific kind of person, and how you will reach them", ""],
           ["Price: how much, how often, and how you get paid (ideally in advance)", ""],
           ["Costs: start-up costs and monthly running costs, with amounts", ""],
           ["Money: how many customers cover your costs? What will you earn per hour?", ""],
           ["Competition: who else does this, and what do customers complain about?", ""],
           ["Biggest risks: the top three, and what you will do about each", ""],
           ["First test: the smallest cheap test, and what result means 'go' or 'stop'", ""],
           ["You: why you can do this, and how many hours a week you can give it", ""]])])

# ---- cap table (formulas, so it recalculates)
wb = Workbook(); ws = wb.active; ws.title = "Cap table"
head = PatternFill("solid", fgColor="DDD0FF"); bold = Font(bold=True)
ws.append(["Shareholder", "Type", "Shares", "% of issued", "% fully diluted"])
rows = [("Founder A", "Ordinary", 5000), ("Founder B", "Ordinary", 5000), ("Option pool (reserved)", "Options", 1000)]
for i, (n, t, s) in enumerate(rows, start=2):
    ws.append([n, t, s, f'=IF(B{i}="Options",0,C{i}/SUMIF($B$2:$B$20,"<>Options",$C$2:$C$20))', f"=C{i}/SUM($C$2:$C$20)"])
for r in range(len(rows) + 2, 21):
    ws.append(["", "", None, f'=IF(C{r}="","",IF(B{r}="Options",0,C{r}/SUMIF($B$2:$B$20,"<>Options",$C$2:$C$20)))', f'=IF(C{r}="","",C{r}/SUM($C$2:$C$20))'])
ws["A22"] = "Total"; ws["C22"] = "=SUM(C2:C20)"; ws["D22"] = "=SUM(D2:D20)"; ws["E22"] = "=SUM(E2:E20)"
for c in ws[1] + ws[22]: c.font = bold
for c in ws[1]: c.fill = head
for r in range(2, 23):
    ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = "0.0%"
for col, w in zip("ABCDE", (26, 12, 12, 14, 16)): ws.column_dimensions[col].width = w
ws["G1"] = "How to use"; ws["G1"].font = bold
for i, t in enumerate(["Type 'Ordinary', 'Preference' or 'Options' in column B.", "Change shares in column C; percentages update.", "Use the Round sheet to see dilution from a new investment.", "A template, not financial or legal advice."], start=2): ws[f"G{i}"] = t
r = wb.create_sheet("Round")
r.append(["New investment round", "Value"]); r["A1"].font = r["B1"].font = bold
for i, (k, v) in enumerate([("Pre-money valuation (£)", 400000), ("Investment (£)", 100000), ("Shares before round", "='Cap table'!C22"), ("Post-money valuation (£)", "=B2+B3"),
                            ("Investor's share of the company", "=B3/B5"), ("New shares issued", "=B4*B6/(1-B6)"), ("Price per share (£)", "=B3/B7"), ("Founders' share before", "=('Cap table'!C2+'Cap table'!C3)/B4"), ("Founders' share after", "=('Cap table'!C2+'Cap table'!C3)/(B4+B7)")], start=2):
    r[f"A{i}"] = k; r[f"B{i}"] = v
for c in ("B6", "B9", "B10"): r[c].number_format = "0.0%"
r.column_dimensions["A"].width = 34; r.column_dimensions["B"].width = 16
wb.save(os.path.join(OUT, "cap-table.xlsx"))

# ---- pitch deck
prs = Presentation(); lay = prs.slide_layouts[1]
slides = [("Title", "Company name. One line on what you do. Your name and contact."), ("Problem", "Who has the problem? How do they cope today? What does it cost them? One real quote or number."),
          ("Solution", "What you built, in one sentence, and one picture or demo."), ("Why now", "What changed that makes this possible or needed today?"),
          ("Market", "Who will buy, how many of them, and how you worked the number out (bottom up)."), ("Product", "How it works in three steps. Screenshots."),
          ("Traction", "Customers, revenue, waiting list, pilots. Growth month by month."), ("Business model", "How you make money: price, who pays, how often."),
          ("Competition", "Who else is there, and why customers choose you."), ("Go to market", "How you find your next 100 customers, and what each costs."),
          ("Team", "Who you are and why you will win. Gaps you plan to fill."), ("Ask", "How much you are raising, what it buys, and what you will have achieved by the next round.")]
for t, b in slides:
    s = prs.slides.add_slide(lay); s.shapes.title.text = t; s.placeholders[1].text = b
prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Keep it to about 12 slides. One idea per slide. Template, not advice."
prs.save(os.path.join(OUT, "pitch-deck-skeleton.pptx"))

# ---- SEIS / EIS checklist (PDF)
styles = getSampleStyleSheet()
pdf = SimpleDocTemplate(os.path.join(OUT, "seis-checklist.pdf"), pagesize=A4, title="SEIS / EIS eligibility checklist", leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50)
def bl(items): return ListFlowable([ListItem(Paragraph("☐ " + i, styles["BodyText"])) for i in items], bulletType="bullet", start="")
story = [Paragraph("SEIS / EIS eligibility checklist (UK)", styles["Title"]),
         Paragraph("Use this to find the obvious problems before you approach investors. The schemes have detailed rules that change, and eligibility is decided by HMRC. Check the current rules on gov.uk and take professional advice before you rely on them. Figures below are the commonly quoted limits and must be confirmed.", styles["Italic"]), Spacer(1, 10),
         Paragraph("The company", styles["Heading2"]),
         bl(["A UK company with a permanent establishment in the UK (a sole trader cannot use these schemes).", "Not listed on a recognised stock exchange, and not controlled by another company.", "SEIS: trading for less than 3 years. EIS: usually less than 7 years since first sale (longer in some cases).",
             "SEIS: fewer than 25 full-time-equivalent employees and gross assets of £350,000 or less before the investment. EIS: fewer than 250 employees and gross assets of £15 million or less.", "A qualifying trade carried on, or being prepared for. Some trades are excluded (for example property, financial services, legal and accounting services, farming, hotels and care homes).",
             "SEIS: has not raised money under EIS or venture capital trusts before. Not previously raised through SEIS beyond its limit.", "The money raised is for growing the trade, not for buying other businesses or shares."]),
         Paragraph("The raise", styles["Heading2"]),
         bl(["SEIS: no more than £250,000 raised in total under SEIS over the company's life.", "EIS: limits on the total raised per year and in a lifetime from venture schemes (check the current figures).", "Shares are new, full-risk ordinary shares with no guaranteed return or protective arrangements.", "Investors do not have, and will not get, a link to the company that stops relief (for example a large shareholding, or being an employee beyond limits)."]),
         Paragraph("Before you start", styles["Heading2"]),
         bl(["Ask HMRC for advance assurance (it takes weeks and many investors expect it).", "Prepare a business plan, accounts and details of the shares for the advance assurance application.", "Keep the money in a UK business bank account and use it as planned.", "File the compliance statement after the shares are issued so investors can claim."]),
         Paragraph("Official help: gov.uk, search 'Seed Enterprise Investment Scheme' and 'Enterprise Investment Scheme'.", styles["BodyText"])]
pdf.build(story)
print("written:", sorted(os.listdir(OUT)))
